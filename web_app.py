import glob
import io
import os
import openpyxl
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="منظومة تدبير الموارد البشرية لمديرية الحوز",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ==================== تنسيقات CSS احترافية ومتجاوبة ====================
st.markdown(
    """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@500;700;800&display=swap');
    
    .stApp {
        font-family: 'Cairo', sans-serif !important;
        direction: rtl !important;
        text-align: right !important;
        background-color: #F8FAFC !important;
    }

    [data-testid="stSidebar"] {
        overflow: hidden !important;
    }
    [data-testid="stSidebar"][aria-expanded="false"] {
        display: none !important;
    }
    
    div[data-testid="InputInstructions"] {
        display: none !important;
    }
    small[data-testid="stFormHelp"] {
        display: none !important;
    }

    .portal-header {
        background: linear-gradient(135deg, #1B365D 0%, #2A5298 100%);
        padding: 18px 24px;
        border-radius: 12px;
        color: white;
        margin-bottom: 18px;
        box-shadow: 0 4px 12px rgba(27, 54, 93, 0.15);
    }
    .portal-title {
        font-size: 21px;
        font-weight: 800;
        margin: 0 0 6px 0;
        color: #FFFFFF;
    }
    .portal-sub {
        font-size: 13px;
        color: #D6E4FF;
        margin: 0;
    }

    .kpi-container {
        display: flex;
        flex-wrap: wrap;
        gap: 10px;
        margin-bottom: 16px;
    }
    .kpi-box {
        flex: 1 1 calc(50% - 10px);
        min-width: 140px;
        background: #FFFFFF;
        padding: 12px 14px;
        border-radius: 10px;
        box-shadow: 0 2px 5px rgba(0, 0, 0, 0.04);
        border-right: 4px solid #2563EB;
        border-top: 1px solid #E2E8F0;
        border-left: 1px solid #E2E8F0;
        border-bottom: 1px solid #E2E8F0;
    }
    .kpi-label {
        color: #64748B;
        font-size: 11px;
        font-weight: 600;
        margin-bottom: 3px;
    }
    .kpi-num {
        color: #0F172A;
        font-size: 16px;
        font-weight: 800;
    }

    @media (max-width: 768px) {
        .kpi-box {
            flex: 1 1 100% !important;
        }
        .portal-title {
            font-size: 18px !important;
        }
        div[data-testid="stDialog"] div[role="dialog"] {
            width: 95vw !important;
            max-width: 95vw !important;
            padding: 12px !important;
        }
    }
</style>
""",
    unsafe_allow_html=True,
)


def clean_value(val):
  if pd.isna(val) or val is None:
    return ""
  s = str(val).strip()
  if s.endswith(".0"):
    s = s[:-2]
  return s


def find_excel_file():
  possible_names = [
      "RH2026.xlsx",
      "rh2026.xlsx",
      "Rh2026.xlsx",
      "RH2026.XLSX",
      "RH.xlsx",
      "rh.xlsx",
  ]
  for name in possible_names:
    if os.path.exists(name):
      return name
  rh_files = glob.glob("*RH*.xlsx") + glob.glob("*rh*.xlsx")
  if rh_files:
    return rh_files[0]
  return None


@st.cache_data
def load_data(source):
  try:
    df = pd.read_excel(source, dtype=str).fillna("")
    df.columns = [str(c).strip() for c in df.columns]

    for c in [
        "PPR",
        "CIN",
        "C_ETAB",
        "CD_ETAB",
        "الوحدة المدرسية",
        "c_unité",
        "c_unite",
    ]:
      if c in df.columns:
        df[c] = df[c].apply(clean_value)

    if "الوحدة المدرسية.1" in df.columns:
      df["اسم_الوحدة_المدرسية"] = (
          df["الوحدة المدرسية.1"].astype(str).str.strip()
      )
    elif "اسم الفرعية" in df.columns:
      df["اسم_الوحدة_المدرسية"] = df["اسم الفرعية"].astype(str).str.strip()
    else:
      df["اسم_الوحدة_المدرسية"] = (
          df.get("الوحدة المدرسية", "").astype(str).str.strip()
      )

    if "c_unité" in df.columns:
      df["كود_الوحدة_المدرسية"] = df["c_unité"].astype(str).str.strip()
    elif "c_unite" in df.columns:
      df["كود_الوحدة_المدرسية"] = df["c_unite"].astype(str).str.strip()
    elif "الوحدة المدرسية" in df.columns:
      df["كود_الوحدة_المدرسية"] = df["الوحدة المدرسية"].astype(str).str.strip()
    else:
      df["كود_الوحدة_المدرسية"] = df.get("C_ETAB", "").astype(str).str.strip()

    return df
  except Exception as e:
    st.error(f"حدث خطأ أثناء قراءة ملف الإكسيل: {e}")
    return pd.DataFrame()


def export_etab_excel_with_header(df_to_export, etab_code, institution_name):
  """تصدير ملف إكسيل يحمل الترويسة الرسمية لوزارة التربية الوطنية في السطور الأولى"""
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "لائحة العاملين"
  ws.views.sheetView[0].rightToLeft = True

  # الترويسة الرسمية
  headers_lines = [
      "المملكة المغربية",
      "وزارة التربية الوطنية والتعليم الأولي والرياضة",
      "الأكاديمية الجهوية للتربية والتكوين جهة مراكش - آسفي",
      "المديرية الإقليمية بالحوز - مصلحة تدبير الموارد البشرية",
      (
          f"لائحة العاملين بمؤسسة: {institution_name} (رمز المؤسسة:"
          f" {etab_code})"
      ),
      "",
  ]

  for r_idx, line in enumerate(headers_lines, start=1):
    cell = ws.cell(row=r_idx, column=1, value=line)
    cell.font = openpyxl.styles.Font(
        name="Arial", size=11, bold=True, color="1B365D"
    )

  # كتابة عناوين الأعمدة
  cols = list(df_to_export.columns)
  for c_idx, col_name in enumerate(cols, start=1):
    c = ws.cell(row=7, column=c_idx, value=col_name)
    c.font = openpyxl.styles.Font(name="Arial", size=10, bold=True, color="FFFFFF")
    c.fill = openpyxl.styles.PatternFill(
        start_color="1F538D", end_color="1F538D", fill_type="solid"
    )
    c.alignment = openpyxl.styles.Alignment(
        horizontal="center", vertical="center"
    )

  # كتابة البيانات
  for r_i, (_, row_data) in enumerate(df_to_export.iterrows(), start=8):
    for c_i, val in enumerate(row_data, start=1):
      cell = ws.cell(row=r_i, column=c_i, value=str(val))
      cell.font = openpyxl.styles.Font(name="Arial", size=9)
      cell.alignment = openpyxl.styles.Alignment(
          horizontal="center", vertical="center"
      )

  out = io.BytesIO()
  wb.save(out)
  return out.getvalue()


def generate_printable_html(df_to_print, etab_code, institution_name):
  """توليد وثيقة A4 مهيأة للطباعة المباشرة مع الترويسة الرسمية"""
  table_html = df_to_print.to_html(index=False, classes="print-table")
  html_doc = f"""
  <!DOCTYPE html>
  <html dir="rtl" lang="ar">
  <head>
      <meta charset="utf-8">
      <title>لائحة العاملين - {institution_name}</title>
      <style>
          @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@600;700;800&display=swap');
          body {{
              font-family: 'Cairo', Tahoma, sans-serif;
              margin: 20px;
              color: #000;
          }}
          .header {{
              text-align: center;
              border-bottom: 2px solid #1B365D;
              padding-bottom: 12px;
              margin-bottom: 20px;
          }}
          .header h3, .header h4, .header h5 {{
              margin: 4px 0;
          }}
          .inst-info {{
              background-color: #f1f5f9;
              padding: 10px 15px;
              border-radius: 6px;
              margin-bottom: 15px;
              font-weight: bold;
              display: flex;
              justify-content: space-between;
          }}
          .print-table {{
              width: 100%;
              border-collapse: collapse;
              margin-top: 10px;
              font-size: 11px;
          }}
          .print-table th, .print-table td {{
              border: 1px solid #334155;
              padding: 6px 8px;
              text-align: center;
          }}
          .print-table th {{
              background-color: #1E3A8A;
              color: white;
          }}
          .print-table tr:nth-child(even) {{
              background-color: #F8FAFC;
          }}
          @media print {{
              @page {{ size: A4 landscape; margin: 12mm; }}
              .no-print {{ display: none; }}
          }}
      </style>
  </head>
  <body>
      <div class="header">
          <h4>المملكة المغربية</h4>
          <h3>وزارة التربية الوطنية والتعليم الأولي والرياضة</h3>
          <h4>الأكاديمية الجهوية للتربية والتكوين جهة مراكش - آسفي</h4>
          <h5>المديرية الإقليمية بالحوز - مصلحة تدبير الموارد البشرية</h5>
      </div>
      <div class="inst-info">
          <span>المؤسسة: {institution_name}</span>
          <span>كود المؤسسة (C_ETAB): {etab_code}</span>
          <span>إجمالي العاملين: {len(df_to_print)}</span>
      </div>
      {table_html}
      <script>
          window.onload = function() {{ window.print(); }}
      </script>
  </body>
  </html>
  """
  return html_doc


detected_file = find_excel_file()
if detected_file:
  df_source = load_data(detected_file)
else:
  st.sidebar.warning("لم يتم العثور على ملف RH2026.xlsx في المستودع.")
  uploaded = st.sidebar.file_uploader(
      "يرجى رفع ملف RH2026.xlsx يدوياً:", type=["xlsx", "xls"]
  )
  if uploaded:
    df_source = load_data(uploaded)
  else:
    df_source = pd.DataFrame()

if df_source.empty:
  st.info("👈 يرجى رفع ملف قاعدة البيانات لبدء التشغيل.")
  st.stop()

cols_all = list(df_source.columns)


def detect_col(candidates, fallback):
  for cand in candidates:
    for c in cols_all:
      if c.strip().lower() == cand.strip().lower():
        return c
  return fallback


COL_C_ETAB = detect_col(["C_ETAB", "كود المؤسسة"], cols_all[0])
COL_UNIT_CODE = "كود_الوحدة_المدرسية"
COL_UNIT_NAME = "اسم_الوحدة_المدرسية"
COL_INSTITUTION = detect_col(
    ["المؤسسة", "اسم المجموعة", "المجموعة المدرسية"], COL_UNIT_NAME
)
COL_TASK = detect_col(
    ["المهام", "FONCTION", "fonct", "المهمة"],
    cols_all[1] if len(cols_all) > 1 else cols_all[0],
)
COL_PPR = detect_col(["PPR", "رقم التأجير", "تأجير"], cols_all[0])

if "df" not in st.session_state or len(st.session_state.df) != len(df_source):
  st.session_state.df = df_source.copy()

df = st.session_state.df

# ==================== نافذة التعديل المنبثقة ====================
if hasattr(st, "dialog"):
  dialog_decorator = st.dialog(
      "✏️ بطاقة تعديل المعطيات الشاملة للأستاذ", width="large"
  )
else:

  def dialog_decorator(func):

    def wrapper(*args, **kwargs):
      with st.expander("✏️ بطاقة تعديل المعطيات الشاملة للأستاذ", expanded=True):
        return func(*args, **kwargs)

    return wrapper


@dialog_decorator
def show_teacher_edit_dialog(ppr_selected):
  match_indices = df.index[df[COL_PPR] == str(ppr_selected)].tolist()
  if not match_indices:
    st.error("تعذر العثور على سجل هذا الأستاذ!")
    return

  idx = match_indices[0]
  t = df.loc[idx]

  st.markdown(
      f"### 👤 {t.get('NOM_AR', '')} | {t.get('NOM_FR', '')} (رقم التأجير:"
      f" `{ppr_selected}`)"
  )

  with st.form("modal_edit_teacher_form"):
    st.markdown("##### 📌 المعطيات الشخصية والمهنية")
    c1, c2 = st.columns(2)
    v_nom_ar = c1.text_input("الاسم الكامل بالعربية:", t.get("NOM_AR", ""))
    v_nom_fr = c2.text_input("الاسم الكامل بالفرنسية:", t.get("NOM_FR", ""))

    c3, c4 = st.columns(2)
    v_cin = c3.text_input("رقم البطاقة الوطنية (CIN):", t.get("CIN", ""))
    v_sex = c4.text_input("الجنس (M/F):", t.get("SEX", ""))

    c5, c6 = st.columns(2)
    v_sit = c5.text_input("الحالة العائلية:", t.get("SIT_FAM", ""))
    v_nais = c6.text_input("تاريخ الازدياد:", t.get("DATE_NAISSANCE", ""))

    st.markdown("##### 🏫 معطيات التعيين والوحدة المدرسية")
    c7, c8 = st.columns(2)
    v_c_etab = c7.text_input("كود المؤسسة (C_ETAB):", t.get(COL_C_ETAB, ""))
    v_u_code = c8.text_input(
        "كود الوحدة المدرسية:", t.get(COL_UNIT_CODE, "")
    )

    c9, c10 = st.columns(2)
    v_u_name = c9.text_input(
        "اسم الوحدة المدرسية (الفرعية):", t.get(COL_UNIT_NAME, "")
    )
    v_inst = c10.text_input(
        "المجموعة المدرسية المركزية:", t.get(COL_INSTITUTION, "")
    )

    c11, c12 = st.columns(2)
    v_task = c11.text_input("المهمة الحالية:", t.get(COL_TASK, ""))
    v_cadre = c12.text_input("الإطار الإداري:", t.get("الاطار", ""))

    st.markdown("##### 📜 الوضعية الإدارية والترقية")
    c13, c14 = st.columns(2)
    v_grade = c13.text_input("الدرجة:", t.get("الدرجة", ""))
    v_d_aff = c14.text_input(
        "تاريخ التعيين بالمؤسسة:", t.get("DATE_DEBUT_AFF", "")
    )

    c15, c16 = st.columns(2)
    v_d_rec = c15.text_input("تاريخ التوظيف:", t.get("DATE_ANC_ADM", ""))
    v_d_gd = c16.text_input("تاريخ مفعول الدرجة:", t.get("DATE_EFFET_GD", ""))

    st.write("")
    submitted = st.form_submit_button(
        "💾 حفظ كافة التعديلات في قاعدة البيانات",
        type="primary",
        use_container_width=True,
    )
    if submitted:
      updates = {
          "NOM_AR": v_nom_ar,
          "NOM_FR": v_nom_fr,
          "CIN": v_cin,
          "SEX": v_sex,
          "SIT_FAM": v_sit,
          "DATE_NAISSANCE": v_nais,
          COL_C_ETAB: v_c_etab,
          COL_UNIT_CODE: v_u_code,
          COL_UNIT_NAME: v_u_name,
          "اسم_الوحدة_المدرسية": v_u_name,
          "كود_الوحدة_المدرسية": v_u_code,
          COL_INSTITUTION: v_inst,
          COL_TASK: v_task,
          "الاطار": v_cadre,
          "الدرجة": v_grade,
          "DATE_DEBUT_AFF": v_d_aff,
          "DATE_ANC_ADM": v_d_rec,
          "DATE_EFFET_GD": v_d_gd,
      }
      for col_name, val in updates.items():
        if col_name in df.columns:
          df.loc[idx, col_name] = clean_value(val)

      st.session_state.df = df
      st.success("✓ تم حفظ كافة التعديلات وتحديث السجل بنجاح!")
      st.rerun()


# ==================== الشريط الجانبي ====================
with st.sidebar:
  st.markdown("### 🏛️ المديرية الإقليمية بالحوز")
  st.metric("إجمالي الموظفين", f"{len(df):,} فرد")

  if st.button("🔄 إعادة ضبط وتحديث القاعدة", use_container_width=True):
    if detected_file:
      st.session_state.df = load_data(detected_file)
      st.success("تم تحديث البيانات بنجاح!")
      st.rerun()

# ==================== الشريط الأزرق العلوي (الترويسة) ====================
st.markdown(
    """
    <div class="portal-header">
        <div>
            <h1 class="portal-title">منظومة تدبير الموارد البشرية لمديرية الحوز</h1>
            <p class="portal-sub">وزارة التربية الوطنية والتعليم الأولي والرياضة - مصلحة تدبير الموارد البشرية 2026</p>
        </div>
    </div>
""",
    unsafe_allow_html=True,
)

# ==================== التبويبات ====================
tab_etab, tab_search, tab_stats, tab_edit = st.tabs([
    "🏫 استعلام المؤسسة والطباعة",
    "🔍 البحث السريع",
    "📊 المؤشرات",
    "✏️ تعديل المعطيات",
])

# -------------------------------------------------------------
# 1. تبويب استعلام المؤسسة والطباعة الرسمية
# -------------------------------------------------------------
with tab_etab:
  with st.form("etab_filter_form"):
    st.markdown("##### 📍 تحديد المؤسسة التعليمية بكود المؤسسة (C_ETAB)")
    col_c1, col_c2 = st.columns([1, 1])

    etab_input = (
        col_c1.text_input(
            "أدخل كود المؤسسة (C_ETAB):", placeholder="مثال: 03254L"
        )
        .strip()
        .upper()
    )

    all_etabs = [""]
    if COL_C_ETAB in df.columns:
      all_etabs += sorted(
          [x for x in df[COL_C_ETAB].unique() if x and x != "nan"]
      )

    etab_dropdown = col_c2.selectbox("أو اختر من القائمة:", all_etabs)

    btn_apply = st.form_submit_button(
        "🔎 تأكيد الاستعلام", use_container_width=True, type="primary"
    )

  target_code = etab_input if etab_input else etab_dropdown

  if target_code:
    res_etab = df[
        df[COL_C_ETAB].astype(str).str.strip().str.upper() == target_code
    ]

    if not res_etab.empty:
      r0 = res_etab.iloc[0]
      institution_name = r0.get(COL_INSTITUTION, "")

      total_staff = len(res_etab)
      females_cnt = (
          (res_etab["SEX"].str.upper() == "F").sum()
          if "SEX" in res_etab.columns
          else 0
      )
      males_cnt = (
          (res_etab["SEX"].str.upper() == "M").sum()
          if "SEX" in res_etab.columns
          else 0
      )

      st.markdown(
          f"""
                <div class="kpi-container">
                    <div class="kpi-box" style="border-right-color:#2563EB;">
                        <div class="kpi-label">المجموعة المدرسية / المؤسسة</div>
                        <div class="kpi-num" style="font-size:16px;">{institution_name if institution_name else 'مؤسسة مستقلة'}</div>
                    </div>
                    <div class="kpi-box" style="border-right-color:#0D9488;">
                        <div class="kpi-label">كود المؤسسة (C_ETAB)</div>
                        <div class="kpi-num">{target_code}</div>
                    </div>
                    <div class="kpi-box" style="border-right-color:#8B5CF6;">
                        <div class="kpi-label">مجموع العاملين</div>
                        <div class="kpi-num">{total_staff} أطر</div>
                    </div>
                    <div class="kpi-box" style="border-right-color:#F59E0B;">
                        <div class="kpi-label">توزيع الجنسين</div>
                        <div class="kpi-num" style="font-size:15px;">{females_cnt} إناث | {males_cnt} ذكور</div>
                    </div>
                </div>
            """,
          unsafe_allow_html=True,
      )

      st.markdown("#### 🏫 الوحدات المدرسية والفرعيات التابعة:")

      sub_units_summary = (
          res_etab.groupby([COL_UNIT_CODE, COL_UNIT_NAME])
          .size()
          .reset_index(name="عدد العاملين")
      )
      sub_units_summary.columns = [
          "كود الوحدة",
          "اسم الوحدة المدرسية (الفرعية)",
          "عدد العاملين",
      ]
      sub_units_summary["النسبة"] = (
          sub_units_summary["عدد العاملين"] / total_staff * 100
      ).map("{:.1f}%".format)

      st.dataframe(sub_units_summary, use_container_width=True, hide_index=True)

      unit_opts = ["عرض جميع الوحدات المدرسية التابعة"]
      for _, r in sub_units_summary.iterrows():
        u_code = r.get("كود الوحدة", "")
        u_name = r.get("اسم الوحدة المدرسية (الفرعية)", "")
        unit_opts.append(f"{u_code} - {u_name}")

      selected_unit_choice = st.selectbox(
          "تصفية اللائحة حسب فرعية معينة:",
          unit_opts,
          key="unit_select_choice_etab",
      )

      if selected_unit_choice != "عرض جميع الوحدات المدرسية التابعة":
        unit_code_selected = (
            selected_unit_choice.split(" - ")[0].strip().upper()
        )
        final_display_df = res_etab[
            res_etab[COL_UNIT_CODE].astype(str).str.strip().str.upper()
            == unit_code_selected
        ]
      else:
        final_display_df = res_etab

      st.markdown("---")
      st.markdown("#### 📋 اللائحة الاسمية للعاملين بالمؤسسة:")
      st.caption("💡 انقر على سطر أي أستاذ بالجدول لفتح بطاقة التعديل.")

      disp_cols = [
          COL_PPR,
          "CIN",
          "NOM_AR",
          "NOM_FR",
          COL_TASK,
          COL_UNIT_NAME,
          COL_UNIT_CODE,
          COL_C_ETAB,
          "الدرجة",
          "DATE_DEBUT_AFF",
      ]
      clean_disp = list(dict.fromkeys(disp_cols))
      avail_cols = [c for c in clean_disp if c in final_display_df.columns]

      rename_disp = {
          COL_PPR: "رقم التأجير",
          "CIN": "رقم البطاقة",
          "NOM_AR": "الاسم بالعربية",
          "NOM_FR": "الاسم بالفرنسية",
          COL_TASK: "المهمة",
          COL_UNIT_NAME: "اسم الوحدة المدرسية",
          COL_UNIT_CODE: "كود الوحدة",
          COL_C_ETAB: "كود المؤسسة",
          "الدرجة": "الدرجة",
          "DATE_DEBUT_AFF": "تاريخ التعيين",
      }

      rendered_df = final_display_df[avail_cols].rename(columns=rename_disp)

      selection = st.dataframe(
          rendered_df,
          use_container_width=True,
          height=300,
          on_select="rerun",
          selection_mode="single-row",
      )

      if (
          selection
          and selection.get("selection")
          and selection["selection"].get("rows")
      ):
        selected_row_idx = selection["selection"]["rows"][0]
        clicked_ppr = str(
            final_display_df.iloc[selected_row_idx][COL_PPR]
        ).strip()
        show_teacher_edit_dialog(clicked_ppr)

      # ==================== أزرار الطباعة الرسمية وتصدير الإكسيل ====================
      st.markdown("---")
      col_btn1, col_btn2 = st.columns(2)

      # 1. زر تصدير Excel رسمي بالترويسة
      excel_official = export_etab_excel_with_header(
          rendered_df, target_code, institution_name
      )
      col_btn1.download_button(
          label="📊 تحميل ملف Excel بالترويسة الرسمية",
          data=excel_official,
          file_name=f"RH_{target_code}_Official.xlsx",
          mime=(
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
          ),
          use_container_width=True,
      )

      # 2. زر الطباعة الفورية A4
      html_print = generate_printable_html(
          rendered_df, target_code, institution_name
      )
      col_btn2.download_button(
          label="🖨️ طباعة اللائحة الرسمية (A4 / PDF)",
          data=html_print,
          file_name=f"Imprimer_{target_code}.html",
          mime="text/html",
          use_container_width=True,
      )

    else:
      st.warning(
          "لم يتم العثور على أي موظف مسجل بكود المؤسسة C_ETAB:"
          f" `{target_code}`"
      )
  else:
    st.info(
        "👈 المرجو إدخال كود المؤسسة (C_ETAB) ثم الضغط على **تأكيد الاستعلام**."
    )

# -------------------------------------------------------------
# 2. تبويب البحث الفردي
# -------------------------------------------------------------
with tab_search:
  with st.form("search_form"):
    st.markdown("##### 🔎 استعلام بالرقم المالي، الاسم أو البطاقة الوطنية")
    query_val = st.text_input(
        "أدخل رقم التأجير (PPR)، الاسم، أو رقم البطاقة (CIN):",
        placeholder="مثال: 1114650 أو محمد",
    ).strip()
    btn_query = st.form_submit_button(
        "🔍 تنفيذ الاستعلام", use_container_width=True, type="primary"
    )

  if query_val:
    q_clean = clean_value(query_val).lower()

    mask_ppr = df[COL_PPR].astype(str).str.strip() == q_clean
    mask_cin = (
        df["CIN"].astype(str).str.strip().str.lower() == q_clean
        if "CIN" in df.columns
        else False
    )
    mask_names = df["NOM_AR"].astype(str).str.contains(
        query_val, regex=False, na=False
    ) | df["NOM_FR"].astype(str).str.lower().str.contains(
        q_clean, regex=False, na=False
    )
    res_q = df[mask_ppr | mask_cin | mask_names]

    if not res_q.empty:
      st.success(
          f"تم العثور على **{len(res_q)}** نتيجة مطابقة (انقر على السطر"
          " لتعديله):"
      )
      raw_disp_cols = [
          COL_PPR,
          "CIN",
          "NOM_AR",
          "NOM_FR",
          COL_TASK,
          "الاطار",
          "الدرجة",
          COL_UNIT_NAME,
          COL_UNIT_CODE,
          COL_C_ETAB,
          "DATE_DEBUT_AFF",
      ]
      disp_cols = list(dict.fromkeys(raw_disp_cols))
      cols_to_render = [c for c in disp_cols if c in res_q.columns]

      rename_disp_q = {
          COL_PPR: "رقم التأجير",
          "CIN": "رقم البطاقة",
          "NOM_AR": "الاسم بالعربية",
          "NOM_FR": "الاسم بالفرنسية",
          COL_TASK: "المهمة",
          COL_UNIT_NAME: "اسم الوحدة المدرسية",
          COL_UNIT_CODE: "كود الوحدة",
          COL_C_ETAB: "كود المؤسسة",
          "الدرجة": "الدرجة",
          "DATE_DEBUT_AFF": "تاريخ التعيين",
      }

      rendered_q = res_q[cols_to_render].rename(columns=rename_disp_q)

      sel_q = st.dataframe(
          rendered_q,
          use_container_width=True,
          on_select="rerun",
          selection_mode="single-row",
      )

      if sel_q and sel_q.get("selection") and sel_q["selection"].get("rows"):
        r_idx = sel_q["selection"]["rows"][0]
        clicked_q_ppr = str(res_q.iloc[r_idx][COL_PPR]).strip()
        show_teacher_edit_dialog(clicked_q_ppr)
    else:
      st.warning("لم يتم العثور على أي موظف بهذه المعطيات.")

# -------------------------------------------------------------
# 3. تبويب لوحة الإحصاء
# -------------------------------------------------------------
with tab_stats:
  st.markdown("##### 📈 مؤشرات التوزيع الإجمالي للموارد البشرية")

  total_all = len(df)
  females_all = (
      (df["SEX"].str.upper() == "F").sum() if "SEX" in df.columns else 0
  )
  males_all = (df["SEX"].str.upper() == "M").sum() if "SEX" in df.columns else 0
  main_etab_count = df[COL_C_ETAB].nunique() if COL_C_ETAB in df.columns else 0

  st.markdown(
      f"""
        <div class="kpi-container">
            <div class="kpi-box" style="border-right-color:#2563EB;">
                <div class="kpi-label">إجمالي الموظفين</div>
                <div class="kpi-num">{total_all:,} موظف</div>
            </div>
            <div class="kpi-box" style="border-right-color:#8B5CF6;">
                <div class="kpi-label">عدد الإناث</div>
                <div class="kpi-num">{females_all:,} ({females_all/total_all*100:.1f}%)</div>
            </div>
            <div class="kpi-box" style="border-right-color:#0D9488;">
                <div class="kpi-label">عدد الذكور</div>
                <div class="kpi-num">{males_all:,} ({males_all/total_all*100:.1f}%)</div>
            </div>
            <div class="kpi-box" style="border-right-color:#EA580C;">
                <div class="kpi-label">المؤسسات المسجلة</div>
                <div class="kpi-num">{main_etab_count} مؤسسة</div>
            </div>
        </div>
    """,
      unsafe_allow_html=True,
  )

  if COL_TASK in df.columns:
    st.markdown("###### 📊 توزيع المهام الرئيسية:")
    top_tasks = df[COL_TASK].value_counts().head(8).reset_index()
    top_tasks.columns = ["المهمة / الإطار", "العدد الإجمالي"]
    st.dataframe(top_tasks, use_container_width=True, hide_index=True)

# -------------------------------------------------------------
# 4. تبويب تعديل سجل موظف بالبحث المباشر
# -------------------------------------------------------------
with tab_edit:
  st.markdown("##### ✏️ تعديل المعطيات الإدارية لموظف")
  target_ppr = clean_value(
      st.text_input(
          "أدخل رقم التأجير (PPR):", placeholder="مثال: 1114650"
      ).strip()
  )

  if target_ppr and COL_PPR in df.columns:
    match_idx = df.index[
        df[COL_PPR].astype(str).str.strip() == target_ppr
    ].tolist()
    if match_idx:
      t_row = df.loc[match_idx[0]]
      st.success(
          f"تم تحميل سجل: **{t_row.get('NOM_AR', '')}**"
          f" ({t_row.get('NOM_FR', '')})"
      )

      with st.form("edit_employee_form_direct"):
        u_nom_ar = st.text_input(
            "الاسم الكامل بالعربية:", t_row.get("NOM_AR", "")
        )
        u_nom_fr = st.text_input("الاسم بالفرنسية:", t_row.get("NOM_FR", ""))
        u_cin = st.text_input(
            "رقم البطاقة الوطنية (CIN):", t_row.get("CIN", "")
        )
        u_task = st.text_input("المهمة الحالية:", t_row.get(COL_TASK, ""))
        u_c_etab = st.text_input(
            "كود المؤسسة (C_ETAB):", t_row.get(COL_C_ETAB, "")
        )
        u_unit_name = st.text_input(
            "اسم الوحدة المدرسية:", t_row.get(COL_UNIT_NAME, "")
        )

        btn_commit = st.form_submit_button(
            "💾 حفظ وتثبيت التعديلات", type="primary"
        )
        if btn_commit:
          if "NOM_AR" in df.columns:
            df.loc[match_idx[0], "NOM_AR"] = u_nom_ar
          if "NOM_FR" in df.columns:
            df.loc[match_idx[0], "NOM_FR"] = u_nom_fr
          if "CIN" in df.columns:
            df.loc[match_idx[0], "CIN"] = u_cin
          if COL_TASK in df.columns:
            df.loc[match_idx[0], COL_TASK] = u_task
          if COL_C_ETAB in df.columns:
            df.loc[match_idx[0], COL_C_ETAB] = u_c_etab
          if COL_UNIT_NAME in df.columns:
            df.loc[match_idx[0], COL_UNIT_NAME] = u_unit_name
          st.session_state.df = df
          st.success("✓ تم حفظ التعديلات بنجاح!")
          st.rerun()
    else:
      st.warning("رقم التأجير غير موجود بقاعدة البيانات.")
